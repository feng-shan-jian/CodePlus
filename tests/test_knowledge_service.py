"""State/identity regressions use real SQLite and OS locks; SDK checks are opt-in."""

from copy import deepcopy
from dataclasses import asdict, replace
from datetime import datetime, timedelta, timezone
import json
import os
from pathlib import Path
import sqlite3
import subprocess
import sys
import time
import asyncio

import pytest

from codeplus.config import KnowledgeConfig
from codeplus.knowledge.metadata import Metadata
from codeplus.knowledge.models import fingerprint, profile
from codeplus.knowledge.service import KnowledgeService


class TinyEmbedding:
    calls = 0
    loaded = False

    @property
    def tokenizer(self):
        self.loaded = True
        return lambda text, **kwargs: {"input_ids": list(range(len(text))),
                                       "offset_mapping": [(i, i + 1) for i in range(len(text))]}

    def encode_documents(self, texts):
        self.calls += 1
        return [[1.0, 0.0, 0.0] for _ in texts]

    def encode_query(self, query):
        return [1.0, 0.0, 0.0]


class MemoryStore:
    def __init__(self):
        self.rows = {}
        self.collections = {}
        self.search_calls = []
        self.search_results = {}

    def check_health(self):
        return "test"

    def ensure_collection(self, name, profile_hash, dimension, *, create=False, indexing=None):
        contract = (profile_hash, dimension, deepcopy(indexing))
        if create and name not in self.collections:
            self.collections[name] = contract
        assert self.collections[name] == contract, "Collection must use its saved profile"

    def upsert_chunks(self, name, rows):
        self.rows.update((row["chunk_id"], row) for row in rows)

    def verify_document(self, name, doc_id, expected):
        assert {r["chunk_id"] for r in expected} == {key for key, r in self.rows.items() if r["doc_id"] == doc_id}

    def delete_document(self, name, doc_id):
        self.rows = {key: r for key, r in self.rows.items() if r["doc_id"] != doc_id}

    def search_dense(self, name, vector, top_k):
        self.search_calls.append(("dense", name, vector, top_k))
        return self.search_results.get("dense", [{"chunk_id": key, "distance": 1.0}
                                                 for key in self.rows])[:top_k]

    def search_bm25(self, name, query, top_k):
        assert self.collections[name][2] is not None
        self.search_calls.append(("bm25", name, query, top_k))
        # A text-matching stub; real ranking/tokenization belongs to the SDK tests.
        return self.search_results.get("bm25", [{"chunk_id": key, "distance": 2.0}
                                                for key, row in self.rows.items()
                                                if query.casefold() in row["text"].casefold()])[:top_k]

    def close(self):
        pass


@pytest.fixture
def service(tmp_path):
    pytest.importorskip("filelock")
    service = KnowledgeService(KnowledgeConfig(enabled=True, data_dir=str(tmp_path / "library"),
                                               embedding_dimension=3, chunk_tokens=64, chunk_overlap=8))
    service.embedding = TinyEmbedding()
    service._store = MemoryStore()
    yield service
    service.close()


@pytest.fixture
def retrieval_case(service, tmp_path):
    kb = service.create("retrieval")["id"]
    for name in ("dense", "shared", "keyword"):
        source = tmp_path / f"{name}.md"
        source.write_text(f"Evidence for {name}.", encoding="utf-8")
        service.import_document(kb, source)
    dense, shared, keyword = service.store.rows
    service.store.search_results = {
        "dense": [{"chunk_id": dense, "distance": 0.9}, {"chunk_id": shared, "distance": 0.7}],
        "bm25": [{"chunk_id": keyword, "distance": 12.0}, {"chunk_id": shared, "distance": 4.0}],
    }
    return kb, (dense, shared, keyword)


@pytest.mark.parametrize("mode,candidates,top_k", [
    ("auto", 1, 3), ("hybrid", 8, 2), ("dense", 8, 2), ("bm25", 8, 2),
])
def test_retrieval_modes_candidates_scores_and_metadata(service, retrieval_case, monkeypatch, mode, candidates, top_k):
    kb, (dense, shared, keyword) = retrieval_case
    before = service.status(kb)
    service.config = replace(service.config, retrieval_mode=mode, retrieval_candidates=candidates, rrf_k=20, top_k=top_k)
    encoded = []
    def encode(query):
        assert mode != "bm25", "BM25 search must not encode the query"
        encoded.append(query)
        return [1.0, 0.0, 0.0]
    monkeypatch.setattr(service.embedding, "encode_query", encode)
    result = service.search(kb, "原始 Query")
    actual = "hybrid" if mode == "auto" else mode
    lanes = ["dense", "bm25"] if actual == "hybrid" else [actual]
    assert [call[0] for call in service.store.search_calls] == lanes
    limit = max(candidates, top_k) if actual == "hybrid" else top_k
    assert all(call[1] == before["collection_name"] and call[3] == limit for call in service.store.search_calls)
    assert encoded == ([] if mode == "bm25" else ["原始 Query"])
    if "bm25" in lanes:
        assert service.store.search_calls[-1][2] == "原始 Query"
    expected = {"dense": [dense, shared], "bm25": [keyword, shared],
                "hybrid": [shared, *sorted([dense, keyword])]}[actual][:top_k]
    assert [hit.chunk_id for hit in result.hits] == expected
    assert len({hit.chunk_id for hit in result.hits}) == len(result.hits)
    scores = {"dense": [0.9, 0.7], "bm25": [12.0, 4.0], "hybrid": [2 / 22, 1 / 21, 1 / 21]}[actual][:top_k]
    assert [hit.score for hit in result.hits] == pytest.approx(scores)
    assert {hit.score_type for hit in result.hits} == {{"dense": "cosine_similarity", "bm25": "bm25", "hybrid": "rrf"}[actual]}
    assert result.query == "原始 Query" and result.kb_id == kb and result.revision == before["revision"]
    assert result.retrieval == {
        "top_k": top_k, "profile_hash": before["profile_hash"],
        "index": {"dense": "FLAT", "bm25": "SPARSE_INVERTED_INDEX", "hybrid": "FLAT+SPARSE_INVERTED_INDEX"}[actual],
        "metric": {"dense": "COSINE", "bm25": "BM25", "hybrid": "RRF"}[actual],
        "requested_mode": mode, "mode": actual, "candidates": limit,
        "rrf_k": 20 if actual == "hybrid" else None,
        "lane_counts": {lane: 2 if lane in lanes else None for lane in ("dense", "bm25")},
    }
    for hit in result.hits:
        saved = service.source(kb, hit.chunk_id)
        assert saved["text"] == hit.text and saved["generation_id"] == hit.generation_id
        assert saved["source_status"] == "current"
    assert service.status(kb) == before


@pytest.mark.parametrize("empty", [("dense",), ("bm25",), ("dense", "bm25")])
def test_hybrid_empty_lanes_are_successful(service, retrieval_case, empty):
    kb, _ = retrieval_case
    for lane in empty:
        service.store.search_results[lane] = []
    result = service.search(kb, "query", 3)
    assert result.retrieval["lane_counts"] == {lane: 0 if lane in empty else 2 for lane in ("dense", "bm25")}
    remaining = next((hits for lane, hits in service.store.search_results.items() if lane not in empty), [])
    assert [hit.chunk_id for hit in result.hits] == [hit["chunk_id"] for hit in remaining]
    assert all(hit.score_type == "rrf" for hit in result.hits)


@pytest.mark.asyncio
@pytest.mark.parametrize("lane", ["dense", "bm25"])
async def test_hybrid_lane_failure_preserves_cause_and_tool_error(service, retrieval_case, monkeypatch, lane):
    from codeplus.knowledge.citations import KnowledgeContext
    from codeplus.tools.knowledge import SearchKnowledge, SearchParams

    kb, _ = retrieval_case
    before = service.status(kb)
    original = RuntimeError("SDK unavailable")
    def fail(*args):
        raise original
    monkeypatch.setattr(service.store, f"search_{lane}", fail)
    with pytest.raises(RuntimeError, match=f"{lane} retrieval failed: SDK unavailable") as error:
        service.search(kb, "query")
    assert error.value.__cause__ is original
    context = KnowledgeContext(service.config)
    context._service = service
    context.bind({"kb_id": kb})
    result = await SearchKnowledge(context).execute(SearchParams(query="query"))
    assert result.is_error
    assert json.loads(result.output) == {"status": "retrieval_failed", "error": f"{lane} retrieval failed: SDK unavailable"}
    assert not context.evidence and context.revision is None
    assert service.status(kb) == before


@pytest.mark.parametrize("mode", ["auto", "dense", "bm25", "hybrid"])
def test_legacy_retrieval_modes_fail_before_queries(service, tmp_path, monkeypatch, mode):
    service.profile.pop("indexing")
    service.profile_hash = fingerprint(service.profile)
    kb = service.create("old dense")["id"]
    source = tmp_path / "old.md"
    source.write_text("Old evidence", encoding="utf-8")
    service.import_document(kb, source)
    service.profile = profile(service.config)
    service.profile_hash = fingerprint(service.profile)
    service.config = replace(service.config, retrieval_mode=mode)
    if mode in ("bm25", "hybrid"):
        monkeypatch.setattr(service.embedding, "encode_query", lambda *_: pytest.fail("Unsupported mode encoded query"))
        monkeypatch.setattr(service, "_collection", lambda *_: pytest.fail("Unsupported mode queried collection"))
        with pytest.raises(ValueError, match=f"does not support {mode}; create a new hybrid base"):
            service.search(kb, "Old")
        assert not service.store.search_calls
    else:
        result = service.search(kb, "Old", 1)
        assert result.retrieval["requested_mode"] == mode and result.retrieval["mode"] == "dense"
        assert result.retrieval["metric"] == "COSINE" and result.retrieval["candidates"] == 1
        assert result.retrieval["rrf_k"] is None and result.retrieval["lane_counts"] == {"dense": 1, "bm25": None}
        assert result.hits[0].score_type == "cosine_similarity"
        assert [call[0] for call in service.store.search_calls] == ["dense"]


@pytest.mark.parametrize("invalid", ["historical", "removed", "not_ready", "other_base", "missing"])
def test_bm25_candidates_require_current_sources(service, retrieval_case, tmp_path, invalid):
    kb, (_, _, keyword) = retrieval_case
    row = service.source(kb, keyword)
    if invalid == "historical":
        source = tmp_path / "keyword.md"
        source.write_text("Replacement evidence", encoding="utf-8")
        service.import_document(kb, source)
    elif invalid == "removed":
        service.remove(kb, row["doc_id"])
    elif invalid == "not_ready":
        with service.metadata.connect() as db:
            db.execute("UPDATE documents SET state='NEEDS_REPAIR' WHERE id=?", (row["doc_id"],))
    elif invalid == "other_base":
        other = service.create("other base")["id"]
        with service.metadata.connect() as db:
            db.execute("UPDATE documents SET kb_id=? WHERE id=?", (other, row["doc_id"]))
    else:
        keyword = "missing-chunk"
    service.store.search_results["dense"] = []
    service.store.search_results["bm25"] = [{"chunk_id": keyword, "distance": 10.0}]
    for mode in ("bm25", "hybrid"):
        service.config = replace(service.config, retrieval_mode=mode)
        with pytest.raises(ValueError, match="no current source in this knowledge base"):
            service.search(kb, "keyword")


def test_hybrid_queries_fusion_and_sources_share_one_lock(service, retrieval_case, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from contextlib import contextmanager
    from filelock import FileLock, Timeout
    from codeplus.knowledge.retrieval import rrf

    kb, _ = retrieval_case
    stages = []
    current_lock = []
    original_locked = service._locked
    @contextmanager
    def locked(*args, **kwargs):
        with original_locked(*args, **kwargs) as bound:
            token = object()
            current_lock.append(token)
            try:
                yield bound
            finally:
                current_lock.pop()
    monkeypatch.setattr(service, "_locked", locked)
    def cannot_write():
        with pytest.raises(Timeout), FileLock(service.root / f"{kb}.lock", timeout=0):
            pytest.fail("Writer entered during retrieval")
    with ThreadPoolExecutor(max_workers=1) as contender:
        def check(stage):
            stages.append((stage, current_lock[-1]))
            contender.submit(cannot_write).result(timeout=5)
        for lane in ("dense", "bm25"):
            original_search = getattr(service.store, f"search_{lane}")
            def search(*args, lane=lane, original=original_search):
                check(lane)
                return original(*args)
            monkeypatch.setattr(service.store, f"search_{lane}", search)
        def fuse(*args):
            check("rrf")
            return rrf(*args)
        monkeypatch.setattr("codeplus.knowledge.service.rrf", fuse)
        original_connect = service.metadata.connect
        @contextmanager
        def connect():
            with original_connect() as db:
                # SQLite trace callbacks swallow exceptions; record and assert below.
                def trace(sql):
                    if sql.startswith("SELECT c.*, d.source_uri FROM chunks"):
                        stages.append(("source", current_lock[-1] if current_lock else None))
                db.set_trace_callback(trace)
                yield db
        monkeypatch.setattr(service.metadata, "connect", connect)
        result = service.search(kb, "keyword", 3)
    assert [stage for stage, _ in stages] == ["dense", "bm25", "rrf", "source", "source", "source"]
    assert len({token for _, token in stages}) == 1 and stages[0][1] is not None
    assert result.revision == service.status(kb)["revision"]


@pytest.mark.parametrize("body", [
    "中文句子必须保持完整，不能从中间截断。" * 90,
    ("重复段落相同。\r\n\r\n" * 100) + "最后一段。\r\n",
    "𠮷" * 1200 + "。\r\n没有标点的超长句也不能丢字。",
    ("短句。" * 20 + "长" * 480 + "。\r\n") * 3,
])
def test_sentence_chunks_preserve_text_positions_and_token_budget(body):
    from codeplus.knowledge.documents import chunk_markdown, parse_markdown

    def tokenizer(value, add_special_tokens=True):
        # Multi-token Unicode character plus two special tokens; no decode path.
        return {"input_ids": [0] * (sum(2 if c == "𠮷" else 1 for c in value)
                                    + (2 if add_special_tokens else 0))}

    text = "# 第一节\r\n\r\n" + body + "\r\n# 第二节\r\n\r\n正文。\r\n"
    chunks = chunk_markdown(text, parse_markdown(text), tokenizer, "doc", "gen", 512, 64)
    assert len(chunks) > 2
    covered = set()
    previous_end = 0
    for chunk in chunks:
        start, end = chunk.source_spans[0].char_start, chunk.source_spans[-1].char_end
        assert chunk.text == text[start:end]
        assert end > previous_end  # Repeated paragraphs must advance to new occurrences.
        assert not text[previous_end:start].strip()
        assert len(tokenizer(chunk.text)["input_ids"]) <= 512
        assert len({tuple(s.heading_path) for s in chunk.source_spans}) == 1
        for span in chunk.source_spans:
            assert span.line_start == text[:span.char_start].count("\n") + 1
            assert span.line_end == text[:span.char_end - 1].count("\n") + 1
            covered.update(range(span.char_start, span.char_end))
        if body.startswith("中文"):
            assert chunk.text.rstrip().endswith("。")
        previous_end = end
    assert all(c.isspace() or i in covered for i, c in enumerate(text))
    assert not text[previous_end:].strip()
    assert any(a.source_spans[-1].char_end > b.source_spans[0].char_start
               for a, b in zip(chunks, chunks[1:]))


def test_previous_chunk_profile_reads_history_but_requires_new_base_for_import(service, tmp_path):
    from codeplus.knowledge.models import fingerprint

    # Short chunks are identical in v1 and v2; persist real v1 profile/generation IDs.
    service.profile.pop("indexing")
    service.profile["chunking"] = "structure-offsets-v1"
    service.profile_hash = fingerprint(service.profile)
    kb = service.create("previous chunker")["id"]
    source = tmp_path / "old.md"
    source.write_bytes("历史内容。\r\n".encode())
    old = service.import_document(kb, source)
    historical = service.search(kb, "历史").hits[0]
    source.write_bytes("当前内容。\r\n".encode())
    service.import_document(kb, source)
    before = service.status(kb)
    reopened = KnowledgeService(service.config)
    reopened.embedding, reopened._store = TinyEmbedding(), service.store
    result = reopened.search(kb, "当前")
    assert result.hits[0].text == "当前内容。\r\n"
    assert result.retrieval["profile_hash"] == before["profile_hash"] != reopened.profile_hash
    assert reopened.source(kb, historical.chunk_id)["source_status"] == "historical"
    assert Path(reopened.source_context(kb, historical.chunk_id)[0]["original_path"]).read_bytes() == "历史内容。\r\n".encode()
    with pytest.raises(ValueError, match="create a new base"):
        reopened.import_document(kb, source)
    assert reopened.status(kb) == before and not reopened.embedding.loaded
    new_kb = reopened.create("new chunker")["id"]
    source.write_bytes("历史内容。\r\n".encode())
    new = reopened.import_document(new_kb, source)
    assert new["generation_id"] != old["generation_id"]
    # Normalization must not waive any other profile mismatch or accept corruption.
    saved_path = reopened.root / kb / "profile.json"
    saved = json.loads(saved_path.read_text(encoding="utf-8"))
    for key, value in (("revision", "other-model"), ("chunking", "unknown-algorithm")):
        changed = {**saved, key: value}
        saved_path.write_text(json.dumps(changed), encoding="utf-8")
        with reopened.metadata.connect() as db:
            db.execute("UPDATE knowledge_bases SET profile_hash=? WHERE id=?", (fingerprint(changed), kb))
        with pytest.raises(ValueError, match="profile mismatch"):
            reopened.search(kb, "test")
    saved_path.write_text(json.dumps(saved), encoding="utf-8")
    with pytest.raises(ValueError, match="profile mismatch"):
        reopened.search(kb, "test")
    assert reopened.source(kb, historical.chunk_id)["text"] == "历史内容。\r\n"
    reopened.close()


@pytest.mark.asyncio
async def test_turn_citations_history_and_reports(service, tmp_path, monkeypatch):
    from codeplus.knowledge.citations import KnowledgeContext
    from codeplus.tools.knowledge import SearchKnowledge, SearchParams, ReadDocument, ReadParams
    from pydantic import ValidationError

    kb = service.create("citations")["id"]
    source = tmp_path / "policy.md"
    source.write_text("# Policy\n\nRefund deadline is 17 days.\n\n" + "Other conditions. " * 12, encoding="utf-8")
    service.import_document(kb, source)
    context = KnowledgeContext(service.config)
    context._service = service
    context.bind({"kb_id": kb, "top_k": 2})
    result = await SearchKnowledge(context).execute(SearchParams(query="deadline", top_k=2))
    payload = json.loads(result.output)
    assert payload["retrieval"]["requested_mode"] == "auto" and payload["retrieval"]["mode"] == "hybrid"
    hit = payload["hits"][0]
    assert hit["score_type"] == "rrf"
    citation = hit["citation_id"]
    preview = await context.preview(citation)
    assert preview["status"] == "current" and preview["filename"] == "policy.md"
    assert "行" in preview["location"] and preview["text"] == service.source(kb, hit["chunk_id"])["text"]
    with pytest.raises(ValidationError):
        SearchParams(query="test", kb_id="other", file_path="anything")
    denied = await ReadDocument(context).execute(ReadParams(citation_id="K:other:unknown"))
    assert denied.is_error
    read = json.loads((await ReadDocument(context).execute(ReadParams(citation_id=citation))).output)
    assert "17 days" in read["source"]["text"]
    assert len(read["neighbours"]) <= 2
    # Exercise pagination with the actual saved text, using a small output budget.
    monkeypatch.setattr("codeplus.knowledge.citations.READ_CHARS", 9)
    first = await context.read(citation)
    assert first["source"]["truncated"] and first["source"]["next_offset"] == 9
    second = await context.read(citation, 9)
    assert first["source"]["text"] + second["source"]["text"] == service.source(kb, hit["chunk_id"])["text"][:18]
    report = await context.report(f"Refund: 17 days. [{citation}]")
    assert '"revision": 1' in report and '"line_start"' in report and '"generation_id"' in report
    with pytest.raises(ValueError, match="Unverified citation"):
        await context.report("wrong [K:other:unknown]")
    with pytest.raises(ValueError, match="not saved"):
        await context.validate_answer("报告已完成")
    source.write_text("Refund deadline is 28 days.", encoding="utf-8")
    service.import_document(kb, source)
    assert "17 days" in service.source_context(kb, hit["chunk_id"])[0]["text"]
    assert (await context.preview(citation))["status"] == "historical"
    with pytest.raises(ValueError, match="corpus changed"):
        await context.search("deadline")
    context.begin_turn()
    await context.search("deadline")
    with pytest.raises(ValueError, match="Unknown citation"):
        await context.read(citation)
    context.bind(None)
    assert (await SearchKnowledge(context).execute(SearchParams(query="test"))).is_error
    service.remove(kb, hit["doc_id"])
    other_kb = service.create("other scope")["id"]
    for binding in ({"kb_id": other_kb}, None):
        context.bind(binding)
        removed = await context.preview(citation)
        assert removed["status"] == "removed" and removed["text"] == preview["text"]
        assert removed["generation_id"] == preview["generation_id"]
        assert not context.evidence and context.revision is None


@pytest.mark.asyncio
async def test_message_local_citation_presentation_and_unavailable_sources(service, tmp_path):
    from codeplus.knowledge.citations import KnowledgeContext, citation_link

    kb = service.create("presentation")["id"]
    refs = []
    for filename in ("[one]&two.md", "second.md"):
        source = tmp_path / filename
        source.write_text("[bold]literal[/bold] <script>alert(1)</script>", encoding="utf-8")
        doc = service.import_document(kb, source)
        with service.metadata.connect() as db:
            chunk = db.execute("SELECT id FROM chunks WHERE doc_id=?", (doc["id"],)).fetchone()[0]
        refs.append(f"K:{kb}:{chunk}")
    context = KnowledgeContext(service.config)
    context._service = service
    first = await context.presentation(f"First [{refs[0]}] again [{refs[0]}]. [External](https://example.com)")
    second = await context.presentation(f"Second [{refs[1]}]")
    assert first["citations"] == [{"citation_id": refs[0], "number": 1, "label": "[one]&two.md · 第1行"}]
    assert second["citations"][0]["number"] == 1 and second["citations"][0]["citation_id"] == refs[1]
    assert first["display_text"].count(f"[\\[1\\]]({citation_link(refs[0])})") == 2
    assert "[External](https://example.com)" in first["display_text"]
    assert (await context.preview(refs[0]))["text"] == "[bold]literal[/bold] <script>alert(1)</script>"
    invalid = await context.presentation("Invalid [K:bad] or missing [K:absent:chunk]")
    assert all("来源不可用" in ref["label"] for ref in invalid["citations"])
    with pytest.raises(ValueError, match="引用格式无效"):
        await context.preview("K:bad")
    with pytest.raises(ValueError, match="not found"):
        await context.preview("K:absent:chunk")


@pytest.mark.asyncio
async def test_agent_knowledge_both_entries_permissions_and_failures(service, tmp_path, monkeypatch):
    from codeplus.agent import Agent, ErrorEvent, StreamText
    from codeplus.conversation import ConversationManager
    from codeplus.knowledge.citations import KnowledgeContext
    from codeplus.permissions import PermissionChecker, PermissionMode
    from codeplus.permissions.dangerous import DangerousCommandDetector
    from codeplus.permissions.rules import RuleEngine, Rule
    from codeplus.permissions.sandbox import PathSandbox
    from codeplus.tools import create_default_registry
    from codeplus.tools.knowledge import SearchKnowledge, ReadDocument
    from codeplus.tools.base import TextDelta, StreamEnd, ToolCallComplete

    kb = service.create("agent")["id"]
    source = tmp_path / "agent.md"
    source.write_text("Refund deadline is 17 days.", encoding="utf-8")
    service.import_document(kb, source)
    context = KnowledgeContext(service.config)
    context._service = service
    context.bind({"kb_id": kb, "top_k": 2})
    registry = create_default_registry()
    registry.register(SearchKnowledge(context))
    registry.register(ReadDocument(context))
    class Client:
        calls = 0
        bad = False
        async def stream(self, conversation, system="", tools=None):
            self.calls += 1
            assert "Refund deadline is 17 days" not in system
            assert "untrusted" in system
            data = json.loads(conversation.history[-1].tool_results[0].content)
            assert data["status"] == "hits"
            assert data["retrieval"]["mode"] == "hybrid" and data["hits"][0]["score_type"] == "rrf"
            ref = data["hits"][0]["citation_id"]
            yield TextDelta("17 days. [" + ("K:bad" if self.bad else ref) + "]")
            yield StreamEnd("end_turn")
    client = Client()
    agent = Agent(client, registry, "openai-compat", work_dir=str(tmp_path))
    agent.knowledge = context
    original_search = service.search
    queries = []
    def search(kb_id, query, top_k):
        queries.append(query)
        return original_search(kb_id, query, top_k)
    monkeypatch.setattr(service, "search", search)
    conv = ConversationManager()
    conv.add_user_message("deadline?")
    conv.add_system_reminder("MCP internal instructions, never a search query")
    events = [event async for event in agent.run(conv)]
    assert queries == ["deadline?"] and any(isinstance(e, StreamText) for e in events)
    assert "17 days" in await agent.run_to_completion("deadline?")
    assert queries == ["deadline?", "deadline?"]
    client.bad = True
    events = [event async for event in agent.run(ConversationManager(history=conv.history[:1]))]
    assert any(isinstance(e, ErrorEvent) for e in events) and not any(isinstance(e, StreamText) for e in events)
    with pytest.raises(ValueError, match="Unverified"):
        await agent.run_to_completion("deadline?")
    rules = RuleEngine(local_rules_path=tmp_path / "permissions.yaml")
    checker = PermissionChecker(DangerousCommandDetector(), PathSandbox(str(tmp_path)), rules, mode=PermissionMode.PLAN)
    agent.permission_checker = checker
    agent.set_permission_mode(PermissionMode.PLAN)
    report_path = tmp_path / "report.md"
    result = await agent._execute_tool_noninteractive(ToolCallComplete("write", "WriteFile", {
        "file_path": str(report_path), "content": "bad [K:invented]"}))
    assert result.is_error and "Permission denied" in result.output and not report_path.exists()
    class ReportClient:
        wrote = False
        async def stream(self, conversation, system="", tools=None):
            if not self.wrote:
                self.wrote = True
                ref = next(iter(context.evidence))
                yield ToolCallComplete("report", "WriteFile", {
                    "file_path": str(report_path), "content": f"17 days [{ref}]"})
                yield StreamEnd("tool_use")
            else:
                yield TextDelta("报告已完成")
                yield StreamEnd("end_turn")
    agent.client = ReportClient()
    with pytest.raises(ValueError, match="not saved.*Permission denied"):
        await agent.run_to_completion("write a report")
    assert not report_path.exists()
    agent.set_permission_mode(PermissionMode.ACCEPT_EDITS)
    agent.client = ReportClient()
    assert await agent.run_to_completion("write a report") == "报告已完成"
    assert '"revision": 1' in report_path.read_text(encoding="utf-8")
    agent.client = client
    rules.append_local_rule(Rule(tool_name="SearchKnowledge", pattern="*", effect="deny"))
    calls = client.calls
    searched = len(queries)
    with pytest.raises(ValueError, match="Permission denied"):
        await agent.run_to_completion("deadline?")
    conv = ConversationManager()
    conv.add_user_message("deadline?")
    assert any(isinstance(e, ErrorEvent) for e in [e async for e in agent.run(conv)])
    assert len(queries) == searched and client.calls == calls
    agent.permission_checker = None
    monkeypatch.setattr(service.store, "search_dense", lambda *args: [])
    context.begin_turn()
    assert (await context.search("no match"))["status"] == "no_hits"
    def unavailable(*args):
        raise RuntimeError("Milvus unavailable")
    monkeypatch.setattr(service.store, "search_bm25", unavailable)
    with pytest.raises(ValueError, match="not no_hits"):
        await agent.run_to_completion("deadline?")


@pytest.mark.asyncio
async def test_knowledge_pilot_binding_background_import_and_history(service, tmp_path, monkeypatch):
    from codeplus.app import CodePlusApp, ChatInput
    from codeplus.config import ProviderConfig
    from codeplus.conversation import Message
    from codeplus.memory.session import SessionMeta
    from codeplus.source_preview import CitationMarkdown, InlineSourcePreview
    from textual.widgets import Static
    import threading

    monkeypatch.chdir(tmp_path)
    provider = ProviderConfig("test", "openai-compat", "http://127.0.0.1:1", "test", "test")
    app = CodePlusApp([provider], enable_fork=False, knowledge_config=service.config)
    app.knowledge._service = service
    external_links = []
    monkeypatch.setattr(app, "open_url", external_links.append)
    messages = []
    show = app.add_system_message
    def record(text):
        messages.append(text)
        show(text)
    monkeypatch.setattr(app, "add_system_message", record)
    source = tmp_path / "path with spaces.md"
    source.write_text("Knowledge pilot answer is 47.", encoding="utf-8")
    started, release = threading.Event(), threading.Event()
    original_import = service.import_document
    def slow_import(*args):
        started.set()
        assert release.wait(10)
        return original_import(*args)
    monkeypatch.setattr(service, "import_document", slow_import)
    async with app.run_test(size=(110, 35)) as pilot:
        async def enter(text, wait=True):
            inp = app.query_one("#chat-input", ChatInput)
            inp.focus()
            inp.load_text(text)
            await pilot.press("enter")
            await pilot.pause()
            if wait and app._knowledge_task:
                await app._knowledge_task
        await enter('/knowledge create "Pilot library"')
        kb = app.knowledge.binding["kb_id"]
        session_id = app.session.session_id
        assert app.session.meta.knowledge_binding["kb_id"] == kb
        await enter(f'/knowledge import "{source}"', wait=False)
        assert await asyncio.to_thread(started.wait, 5)
        await enter('/knowledge off', wait=False)
        assert app.knowledge.binding["kb_id"] == kb and any("正在执行" in m for m in messages)
        # The real input widget still accepts keystrokes while import is held in a thread.
        await pilot.press("a", "b", "c")
        assert app.query_one("#chat-input", ChatInput).text == "abc"
        app.query_one("#chat-input", ChatInput).clear()
        release.set()
        await app._knowledge_task
        assert any("导入完成：已处理 1/1，成功 1，未变化 0，失败 0" in m for m in messages)
        from codeplus.tools.base import TextDelta, StreamEnd
        class Client:
            async def stream(self, conversation, system="", tools=None):
                ref = next(iter(app.knowledge.evidence))
                yield TextDelta(f"47 [{ref}]\n\n[External](https://example.com)")
                yield StreamEnd("end_turn")
        app.agent.client = Client()
        app.agent.memory_manager = None
        await enter('pilot answer?')
        if app._agent_task:
            await app._agent_task
        assert app.conversation.history[-1].content.startswith("47 [K:")
        await pilot.pause()
        answer = app.query_one(CitationMarkdown)
        assert answer.references[0]["number"] == 1
        # Click the rendered inline [1], through Textual's actual link event path.
        paragraph = answer.query("MarkdownParagraph").first()
        paragraph.scroll_visible(animate=False)
        await pilot.pause()
        await pilot.click(paragraph, offset=(4, 0))
        await pilot.pause()
        preview = app.query_one(InlineSourcePreview)
        assert "Knowledge pilot answer is 47." in str(preview.query_one(".source-text", Static).render())
        assert "当前版本" in str(preview.query_one(".source-details", Static).render())
        await pilot.press("escape")
        await pilot.pause()
        assert not app.query(InlineSourcePreview)
        external = list(answer.query("MarkdownParagraph"))[1]
        external.scroll_visible(animate=False)
        await pilot.pause()
        await pilot.click(external, offset=(2, 0))
        assert external_links == ["https://example.com"]
        await enter('/knowledge sources')
        await enter('/knowledge status')
        hit = service.search(kb, "answer").hits[0]
        ref = f"K:{kb}:{hit.chunk_id}"
        second_file = tmp_path / "second [literal].md"
        second_file.write_text("[bold]Second answer is 63.[/bold]", encoding="utf-8")
        second_doc = service.import_document(kb, second_file)
        second_hit = next(hit for hit in service.search(kb, "answer", 10).hits if hit.doc_id == second_doc["id"])
        second_ref = f"K:{kb}:{second_hit.chunk_id}"
        app.session.append(Message("assistant", f"63 [{second_ref}]. Compare [{ref}]"))
        app.conversation.add_assistant_message(f"63 [{second_ref}]. Compare [{ref}]")
        await enter('/session new')
        assert app.knowledge.binding is None and not app.registry.is_enabled("SearchKnowledge")
        await enter(f'/session resume {session_id}')
        assert app.knowledge.binding["kb_id"] == kb
        await enter('/knowledge off')
        assert not app.conversation.history and app.session.meta.knowledge_binding is None
        await enter(f'/knowledge open {ref}')
        assert any('"historical_sources"' in m and "47" in m for m in messages)
        await enter(f'/session resume {session_id}')
        assert app.knowledge.binding is None and all(ref not in m.content for m in app.conversation.history)
        restored_answers = list(app.query(CitationMarkdown))
        assert len(restored_answers) == 2  # Two persisted answers; no synthetic boundary/keep duplicates.
        assert [answer.references[0]["number"] for answer in restored_answers] == [1, 1]
        assert [answer.references[0]["citation_id"] for answer in restored_answers] == [ref, second_ref]
        # Start in the real input, navigate to the second answer and open its own [1].
        await pilot.press("tab", "tab", "enter")
        await pilot.pause()
        preview = app.query_one(InlineSourcePreview)
        assert "[bold]Second answer is 63.[/bold]" in str(preview.query_one(".source-text", Static).render())
        assert "second [literal].md" in str(preview.query_one(".source-details", Static).render())
        await pilot.press("down", "enter")
        await pilot.pause()
        assert "Knowledge pilot answer is 47." in str(preview.query_one(".source-text", Static).render())
        await pilot.press("escape")
        await pilot.pause()
        assert app.focused is restored_answers[1]
        # A different library and then off leave the original link resolvable on resume.
        await enter('/knowledge create "Second library"')
        assert app.knowledge.binding["kb_id"] != kb
        source.write_text("Updated answer is 99.", encoding="utf-8")
        updated = service.import_document(kb, source)
        service.remove(kb, updated["id"])
        await enter('/knowledge off')
        await enter(f'/session resume {session_id}')
        answer = app.query(CitationMarkdown).first()
        await pilot.press("tab", "enter")
        await pilot.pause()
        preview = app.query_one(InlineSourcePreview)
        assert "已删除" in str(preview.query_one(".source-details", Static).render())
        assert "Knowledge pilot answer is 47." in str(preview.query_one(".source-text", Static).render())
        await pilot.press("escape")
        await pilot.pause()
        assert app.focused is answer and not app.knowledge.evidence
        assert all(ref not in m.content for m in app.conversation.history)
        await enter(f'/knowledge use {kb}')
        await enter('/clear')
        assert app.knowledge.binding is None and not app.registry.is_enabled("ReadDocument")
        # Old metadata defaults off; missing restored library remains bound and blocks answers.
        meta_path = app.session._sessions_dir / f"{app.session.session_id}.meta"
        data = json.loads(meta_path.read_text(encoding="utf-8"))
        data.pop("knowledge_binding")
        meta_path.write_text(json.dumps(data), encoding="utf-8")
        assert SessionMeta.load(meta_path).knowledge_binding is None
        app.session.meta.knowledge_binding = {"kb_id": "f" * 32, "top_k": 2}
        app.session.meta.save(meta_path)
        await enter(f'/session resume {app.session.session_id}')
        assert any("不可用" in m for m in messages)
        assert app.knowledge.binding["kb_id"] == "f" * 32
        app.session.close()


def test_prompt_knowledge_json_exit_and_permissions(service, tmp_path, monkeypatch, capsys):
    from codeplus import __main__ as cli
    from codeplus.config import AppConfig, ProviderConfig
    from codeplus.tools.base import TextDelta, StreamEnd, ToolCallComplete

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(cli.crashlog, "install", lambda: None)
    monkeypatch.setattr(cli.logging, "basicConfig", lambda **kwargs: None)
    kb = service.create("prompt")["id"]
    source = tmp_path / "policy.md"
    source.write_text("Refund deadline is 17 days.", encoding="utf-8")
    service.import_document(kb, source)
    config = AppConfig([ProviderConfig("test", "openai-compat", "http://127.0.0.1:1", "test")],
                       knowledge=service.config, enable_fork=False)
    monkeypatch.setattr(cli, "load_config", lambda: config)
    def new_service(cfg):
        # Each CLI invocation owns its connection/lifecycle, with the same test corpus.
        current = KnowledgeService(cfg)
        current.embedding, current._store = service.embedding, service.store
        current.search = noisy_search
        prepare = current.prepare
        def noisy_prepare(progress):
            print("PREPARE_DIAGNOSTIC")
            return prepare(progress)
        current.prepare = noisy_prepare
        return current
    monkeypatch.setattr("codeplus.knowledge.service.KnowledgeService", new_service)
    async def resolve(provider):
        pass
    monkeypatch.setattr("codeplus.client.resolve_context_window", resolve)
    class Client:
        report = False
        async def stream(self, conversation, system="", tools=None):
            results = [m.tool_results for m in conversation.history if m.tool_results]
            if not results:
                assert not any(t["name"] == "SearchKnowledge" for t in tools)
                yield TextDelta("PLAIN_OK")
            elif self.report and len(results) == 1:
                yield ToolCallComplete("report", "WriteFile", {
                    "file_path": str(tmp_path / "report.md"), "content": "Unapproved report"})
            else:
                ref = json.loads(results[0][0].content)["hits"][0]["citation_id"]
                yield TextDelta(f"17 days [{ref}]")
            yield StreamEnd("end_turn")
    client = Client()
    monkeypatch.setattr("codeplus.client.create_client", lambda provider: client)
    original = service.search
    def noisy_search(*args):
        print("MODEL_LOAD_DIAGNOSTIC")
        return original(*args)
    monkeypatch.setattr(service, "search", noisy_search)
    def run(extra, failure=False):
        monkeypatch.setattr(sys, "argv", ["codeplus", "-p", "deadline?", *extra])
        if failure:
            with pytest.raises(SystemExit) as exit:
                cli.main()
            assert exit.value.code == 1
        else:
            cli.main()
        return capsys.readouterr()
    json_args = ["--knowledge", kb, "--output-format", "stream-json"]
    output = run(json_args)
    events = [json.loads(line) for line in output.out.splitlines()]
    assert events[0]["tool_name"] == events[1]["tool_name"] == "SearchKnowledge"
    assert events[0]["tool_id"] == events[1]["tool_id"]
    assert events[0]["args"]["query"] == "deadline?"
    assert events[-1]["type"] == "result" and "17 days [K:" in events[-1]["result"]
    assert "MODEL_LOAD_DIAGNOSTIC" in output.err
    assert "PREPARE_DIAGNOSTIC" in output.err and "已就绪" in output.err
    assert run(["--knowledge", kb]).out.startswith("17 days [K:")
    assert run([]).out == "PLAIN_OK"
    for bad in ("missing", "../missing"):
        events = [json.loads(line) for line in run(["--knowledge", bad, "--output-format", "stream-json"], True).out.splitlines()]
        assert events[-1]["type"] == "error" and all(e["type"] != "result" for e in events)
    with monkeypatch.context() as patch:
        patch.setattr(config.knowledge, "enabled", False)
        assert "disabled" in run(["--knowledge", kb], True).err
    for effect in ("deny", "ask"):
        rules = tmp_path / ".codeplus" / "permissions.local.yaml"
        rules.write_text(f"- rule: SearchKnowledge(*)\n  effect: {effect}\n", encoding="utf-8")
        events = [json.loads(line) for line in run(json_args, True).out.splitlines()]
        assert events[-1]["type"] == "error"
        assert ("Permission denied" if effect == "deny" else "user rejected") in events[-1]["message"]
    rules.write_text("[]", encoding="utf-8")
    client.report = True
    events = [json.loads(line) for line in run(json_args, True).out.splitlines()]
    assert events[-1]["type"] == "error" and "not saved" in events[-1]["message"]
    assert not (tmp_path / "report.md").exists()


@pytest.mark.asyncio
async def test_remote_knowledge_websocket_scope_progress_and_resume(service, tmp_path, monkeypatch):
    import threading
    import websockets
    from codeplus.config import AppConfig, ProviderConfig
    from codeplus.remote import RemoteServer
    from codeplus.tools.base import TextDelta, StreamEnd

    monkeypatch.chdir(tmp_path)
    config = AppConfig([ProviderConfig("test", "openai-compat", "http://127.0.0.1:1", "test", "test")],
                       knowledge=service.config, enable_fork=False)
    server = RemoteServer(config.providers, config=config)
    server.knowledge._service = service
    server._init_agent()
    started, release = threading.Event(), threading.Event()
    original_import = service.import_document
    def slow_import(*args):
        started.set()
        assert release.wait(10)
        return original_import(*args)
    monkeypatch.setattr(service, "import_document", slow_import)
    answer_started, answer_release = asyncio.Event(), asyncio.Event()
    class Client:
        async def stream(self, conversation, system="", tools=None):
            answer_started.set()
            await answer_release.wait()
            if server.knowledge.binding:
                ref = next(iter(server.knowledge.evidence))
                yield TextDelta(f"17 days [{ref}]")
            else:
                assert not any("K:" in m.content or m.tool_results for m in conversation.history)
                yield TextDelta("OFF_OK")
            yield StreamEnd("end_turn")
    server.agent.client = Client()
    source = tmp_path / "server path with spaces.md"
    source.write_text("Refund deadline is 17 days.", encoding="utf-8")
    try:
        async with websockets.serve(server._ws_handler, "127.0.0.1", 0) as listener:
            port = listener.sockets[0].getsockname()[1]
            async with websockets.connect(f"ws://127.0.0.1:{port}/ws") as ws:
                async def until(kind, ready=lambda events: True):
                    events = []
                    async with asyncio.timeout(10):
                        while True:
                            events.append(json.loads(await ws.recv()))
                            if (kind is None or events[-1]["type"] == kind) and ready(events):
                                return events
                async def send(content):
                    await ws.send(json.dumps({"type": "user_message", "data": {"content": content}}))
                async def command(content):
                    await send(content)
                    events = await until("command_done")
                    assert not any(e["type"] == "error" for e in events), events
                    return events
                await until("commands")
                await command('/knowledge create "remote"')
                kb = server.knowledge.binding["kb_id"]
                session_id = server.session_id
                await send(f'/knowledge import "{source}"')
                assert await asyncio.to_thread(started.wait, 5)
                await send("/knowledge off")
                await send("question during import")
                await ws.send(json.dumps({"type": "ping"}))
                events = await until(None, lambda es: any(e["type"] == "pong" for e in es)
                                     and sum("正在执行" in str(e) for e in es) == 2)
                assert sum("正在执行" in e.get("data", {}).get("message", "") for e in events if isinstance(e.get("data"), dict)) == 2
                assert not answer_started.is_set() and server.knowledge.binding["kb_id"] == kb
                release.set()
                events += await until("command_done")
                assert any("服务器本地" in str(e) for e in events)
                assert any("导入完成：已处理 1/1，成功 1，未变化 0，失败 0" in str(e) for e in events)
                errors = await command(f'/knowledge import "{tmp_path / "missing.md"}"')
                assert any("导入失败" in str(e) for e in errors)
                await command(f"/knowledge use {kb}")
                await send("deadline?")
                await asyncio.wait_for(answer_started.wait(), 5)
                await send("/knowledge off")
                await ws.send(json.dumps({"type": "ping"}))
                events = await until(None, lambda es: any(e["type"] == "pong" for e in es)
                                     and any("正在执行" in str(e) for e in es))
                assert any("正在执行" in str(e) for e in events)
                assert server.knowledge.binding["kb_id"] == kb
                answer_release.set()
                events += await until("loop_complete")
                use = next(e["data"] for e in events if e["type"] == "tool_use")
                result = next(e["data"] for e in events if e["type"] == "tool_result")
                assert use["toolId"] == result["toolId"] and use["toolName"] == "SearchKnowledge"
                ref = next(iter(server.knowledge.evidence))
                rendered = next(e["data"] for e in events if e["type"] == "stream_end" and "17 days" in e["data"]["text"])
                assert rendered["citations"][0]["citation_id"] == ref and "[\\[1\\]]" in rendered["display_text"]
                await command("/session new")
                assert server.knowledge.binding is None
                restored = await command(f"/session resume {session_id}")
                assert server.knowledge.binding["kb_id"] == kb
                assert any(e["type"] == "replay_assistant" and ref in str(e) for e in restored)
                await command("/knowledge off")
                opened = await command(f"/knowledge open {ref}")
                assert any("historical_sources" in str(e) and "17 days" in str(e) for e in opened)
                replay = await command(f"/session resume {session_id}")
                assert server.knowledge.binding is None and not server.knowledge.evidence
                assert any(e["type"] == "replay_assistant" and e["data"]["citations"][0]["citation_id"] == ref for e in replay)
                await ws.send(json.dumps({"type": "source_request", "data": {"citation_id": ref, "request_id": "open-old"}}))
                opened = (await until("source_preview"))[-1]["data"]
                assert opened["request_id"] == "open-old" and "17 days" in opened["source"]["text"]
                await ws.send(json.dumps({"type": "source_request", "data": {"citation_id": "K:bad", "request_id": "invalid"}}))
                assert "引用格式无效" in (await until("source_preview"))[-1]["data"]["error"]
                assert not server.knowledge.evidence and all(ref not in m.content for m in server.conversation.history)
                await send("normal answer")
                assert any("OFF_OK" in str(e) for e in await until("loop_complete"))
                await command(f"/knowledge use {kb}")
                await command("/clear")
                assert server.knowledge.binding is None and server.session_id != session_id
                assert not server.registry.is_enabled("ReadDocument")
                # Existing PROMPT commands must hand their internally generated prompt to the Agent.
                from codeplus.commands.handlers.review import REVIEW_COMMAND
                server.command_registry.register_sync(REVIEW_COMMAND)
                await send("/review")
                assert any("OFF_OK" in str(e) for e in await until("loop_complete"))
                assert any("git diff" in m.content for m in server.conversation.history if m.role == "user")
    finally:
        release.set()
        answer_release.set()
        await asyncio.gather(*server._message_tasks, return_exceptions=True)
        await server._flush_ui_messages()
        server.session.close()


def test_import_identity_sources_and_preparation_failure(service, tmp_path, monkeypatch):
    kb = service.create("sources")["id"]
    source = tmp_path / "source with spaces.md"
    data = "# 标题\r\n\r\n需要准确定位的中文段落。\r\n\r\n```python\r\nprint('hello')\r\n```\r\n".encode()
    source.write_bytes(data)
    doc = service.import_document(kb, source)
    saved = Path(doc["original_path"])
    assert saved.read_bytes() == data and not saved.stat().st_mode & 0o200
    before = service.status(kb)
    assert service.import_document(kb, source)["unchanged"]
    assert service.status(kb) == before and service.embedding.calls == 1
    other = tmp_path / "other.md"
    other.write_bytes(data)
    duplicate_bytes = service.import_document(kb, other)
    assert duplicate_bytes["id"] != doc["id"] and duplicate_bytes["generation_id"] == doc["generation_id"]
    wrong = KnowledgeService(replace(service.config, embedding_revision="same-dimension-other-model"))
    with pytest.raises(ValueError, match="profile mismatch"):
        wrong.import_document(kb, source)
    assert wrong.embedding._model is None and wrong._store is None
    hit = service.search(kb, "中文").hits[0]
    original = saved.read_bytes().decode("utf-8-sig")
    for span in hit.source_spans:
        fragment = original[span.char_start:span.char_end]
        assert fragment in hit.text
        assert fragment in "".join(original.splitlines(keepends=True)[span.line_start-1:span.line_end])
    rows_b = {key: row for key, row in service.store.rows.items() if row["doc_id"] == duplicate_bytes["id"]}
    source.write_text("Changed", encoding="utf-8")
    updated = service.import_document(kb, source)
    assert updated["id"] == doc["id"] and updated["generation_id"] != doc["generation_id"]
    assert service.status(kb)["revision"] == 3 and service.status(kb)["collection_name"] == before["collection_name"]
    assert all(h.generation_id == updated["generation_id"] for h in service.search(kb, "test").hits if h.doc_id == doc["id"])
    assert {key: row for key, row in service.store.rows.items() if row["doc_id"] == duplicate_bytes["id"]} == rows_b
    assert Path(service.source(kb, hit.chunk_id)["original_path"]).read_bytes() == data == saved.read_bytes()
    assert service.remove(kb, doc["id"])["removed"]
    state = service.status(kb)
    assert service.remove(kb, doc["id"])["unchanged"] and service.status(kb) == state
    assert all(h.doc_id != doc["id"] for h in service.search(kb, "test").hits)
    assert service.source(kb, hit.chunk_id)["original_path"] == str(saved)
    source.write_bytes(data)  # Restore a historical generation: stable IDs, retained original mapping.
    restored = service.import_document(kb, source)
    assert restored["generation_id"] == doc["generation_id"] and restored["original_path"] == str(saved)
    assert not restored["removed"] and service.status(kb)["revision"] == 5
    assert service.retry(kb)["unchanged"] and service.status(kb)["revision"] == 5
    # A concurrent replacement wins while the first import prepares its vectors.
    source.write_text("First candidate", encoding="utf-8")
    encode = service.embedding.encode_documents
    def concurrent_update(texts):
        other_writer = KnowledgeService(service.config)
        other_writer.embedding, other_writer._store = TinyEmbedding(), service.store
        source.write_text("Concurrent winner", encoding="utf-8")
        other_writer.import_document(kb, source)
        return encode(texts)
    with monkeypatch.context() as patch:
        patch.setattr(service.embedding, "encode_documents", concurrent_update)
        with pytest.raises(ValueError, match="changed during preparation"):
            service.import_document(kb, source)
    assert Path(service.metadata.document(doc["id"])["original_path"]).read_text() == "Concurrent winner"
    before = service.status(kb)
    fresh = tmp_path / "failure.md"
    fresh.write_text("# New\n\nFailed preparation", encoding="utf-8")
    def fail_model(texts):
        raise RuntimeError("model failed")
    monkeypatch.setattr(service.embedding, "encode_documents", fail_model)
    with pytest.raises(RuntimeError, match="model failed"):
        service.import_document(kb, fresh)
    monkeypatch.setattr(service.embedding, "encode_query", fail_model)
    with pytest.raises(RuntimeError, match="model failed"):
        service.search(kb, "query model failure")
    assert service.status(kb) == before and not list(service.root.glob("preparing-*"))
    empty = service.create("empty")["id"]
    with pytest.raises(ValueError, match="empty"):
        service.search(empty, "test")
    fresh.write_text("\n\n", encoding="utf-8")
    with pytest.raises(ValueError, match="no content"):
        service.import_document(empty, fresh)
    assert service.status(empty)["revision"] == 0


def test_legacy_dense_writes_keep_profile_and_generation(service, tmp_path):
    service.profile.pop("indexing")
    service.profile_hash = fingerprint(service.profile)
    saved_profile = deepcopy(service.profile)
    kb = service.create("legacy dense")["id"]
    reopened = KnowledgeService(service.config)
    reopened.embedding, reopened._store = TinyEmbedding(), service.store
    source = tmp_path / "legacy.md"
    source.write_text("Original contents.", encoding="utf-8")
    original = reopened.import_document(kb, source)
    hit = reopened.search(kb, "contents").hits[0]
    assert reopened.import_document(kb, source)["unchanged"]
    hybrid = reopened.create("hybrid")["id"]
    same = reopened.import_document(hybrid, source)
    assert same["generation_id"] == original["generation_id"]
    source.write_text("Updated contents.", encoding="utf-8")
    changed = reopened.import_document(kb, source)
    assert changed["generation_id"] != original["generation_id"]
    assert reopened.source(kb, hit.chunk_id)["source_status"] == "historical"
    assert reopened.remove(kb, original["id"])["removed"]
    assert reopened.retry(kb)["unchanged"]
    assert reopened.status(kb)["revision"] == 3
    assert reopened.status(kb)["profile_hash"] == fingerprint(saved_profile)
    assert json.loads((reopened.root / kb / "profile.json").read_text(encoding="utf-8")) == saved_profile
    assert service.store.collections[reopened.status(kb)["collection_name"]][2] is None
    assert service.store.collections[reopened.status(hybrid)["collection_name"]][2] == reopened.profile["indexing"]


@pytest.mark.parametrize("change", ["version", "analyzer", "index", "partial", "null", "chunker", "hash"])
def test_unknown_or_corrupt_index_profiles_are_rejected(service, change):
    kb = service.create("invalid profile")["id"]
    saved = deepcopy(service.profile)
    if change == "version":
        saved["indexing"]["version"] = "unknown"
    elif change == "analyzer":
        saved["indexing"]["analyzer"]["tokenizer"]["hmm"] = True
    elif change == "index":
        saved["indexing"]["sparse"]["params"]["bm25_b"] = 0.5
    elif change == "partial":
        del saved["indexing"]["dense"]
    elif change == "null":
        saved["indexing"] = None
    elif change == "chunker":
        saved["chunking"] = "structure-offsets-v1"  # Never a known hybrid profile.
    else:
        saved.pop("indexing")  # Otherwise valid legacy shape with a mismatched saved hash.
    (service.root / kb / "profile.json").write_text(json.dumps(saved), encoding="utf-8")
    if change != "hash":
        with service.metadata.connect() as db:
            db.execute("UPDATE knowledge_bases SET profile_hash=? WHERE id=?", (fingerprint(saved), kb))
    with pytest.raises(ValueError, match="profile mismatch"):
        service.retry(kb)


@pytest.mark.parametrize("failure,legacy", [("upsert", None), ("verify", None),
                                          ("metadata", None), ("upsert", "dense"), ("upsert", "structure")])
def test_failed_commit_retains_pending_and_blocks_search(service, tmp_path, monkeypatch, failure, legacy):
    if legacy:
        service.profile.pop("indexing")
        if legacy == "structure":
            service.profile["chunking"] = "structure-offsets-v1"
        service.profile_hash = fingerprint(service.profile)
    bound_profile = deepcopy(service.profile)
    kb = service.create("failure")["id"]
    if legacy == "dense":
        # An upgraded service still writes the original dense library's profile.
        service.profile = profile(service.config)
        service.profile_hash = fingerprint(service.profile)
    source = tmp_path / "document.md"
    source.write_text("# Long\n\n" + "中文正文 " * 40, encoding="utf-8")
    def fail(*args):
        raise RuntimeError(failure)
    if failure == "upsert":
        monkeypatch.setattr(service.store, "upsert_chunks", fail)
    elif failure == "verify":
        monkeypatch.setattr(service.store, "verify_document", fail)
    else:
        with service.metadata.connect() as db:
            db.execute("CREATE TRIGGER fail_second BEFORE INSERT ON chunks WHEN NEW.ordinal=1 "
                       "BEGIN SELECT RAISE(ABORT, 'metadata'); END")
    with pytest.raises((RuntimeError, sqlite3.IntegrityError), match=failure):
        service.import_document(kb, source)
    state = service.status(kb)
    assert state["state"] == "NEEDS_REPAIR" and state["revision"] == 0
    document, = state["documents"]
    assert document["state"] == "NEEDS_REPAIR" and document["chunk_count"] == 0
    assert document["updated_at"] is None
    assert document["pending_operation"] == "import"
    pending = json.loads(Path(document["pending_path"]).read_text(encoding="utf-8"))
    assert pending["profile"] == bound_profile
    assert fingerprint(pending["profile"]) == state["profile_hash"]
    assert pending["document"]["content_hash"] == document["content_hash"]
    assert len(pending["chunks"]) == len(pending["rows"]) > 1
    with pytest.raises(ValueError, match="NEEDS_REPAIR"):
        service.search(kb, "test")
    with service.metadata.connect() as db:
        db.execute("ALTER TABLE documents DROP COLUMN updated_at")
        db.execute("PRAGMA user_version = 2")
    reopened = KnowledgeService(service.config)
    assert reopened.status(kb) == state
    assert reopened.list_libraries()[0]["document_count"] == 1
    assert reopened.list_libraries()[0]["updated_at"] is None
    with pytest.raises(ValueError, match="NEEDS_REPAIR"):
        reopened.search(kb, "test")
    monkeypatch.undo()
    if failure == "metadata":
        with service.metadata.connect() as db:
            db.execute("DROP TRIGGER fail_second")
    reopened._store = service.store
    reopened.embedding = TinyEmbedding()
    if failure == "upsert":
        path = Path(document["pending_path"])
        corrupt = deepcopy(pending)
        corrupt["profile"]["revision"] = "different-model"
        path.write_text(json.dumps(corrupt), encoding="utf-8")
        with pytest.raises(ValueError, match="registered document/profile"):
            reopened.retry(kb)
        assert reopened.status(kb)["revision"] == 0
        path.write_text(json.dumps(pending), encoding="utf-8")
    assert reopened.retry(kb)["chunk_count"] == len(pending["chunks"])
    assert datetime.fromisoformat(reopened.metadata.document(document["id"])["updated_at"]).tzinfo == timezone.utc
    assert reopened.status(kb)["state"] == "READY" and reopened.status(kb)["revision"] == 1
    assert reopened.retry(kb)["unchanged"]
    assert reopened.search(kb, "test").hits
    if legacy == "structure":
        with pytest.raises(ValueError, match="create a new base"):
            reopened.import_document(kb, source)


@pytest.mark.parametrize("version", [1, 2])
def test_schema_constraints_and_version(service, tmp_path, monkeypatch, version):
    metadata = service.metadata
    with pytest.raises(sqlite3.IntegrityError):
        with metadata.connect() as db:
            db.execute("INSERT INTO chunks VALUES ('c','unknown','g',0,'text','[]','original.md')")
    kb = service.create("S2")["id"]
    source = tmp_path / "s2.md"
    source.write_text("S2 existing document", encoding="utf-8")
    document = service.import_document(kb, source)
    hit = service.search(kb, "test").hits[0]
    with monkeypatch.context() as patch:
        def unavailable(*args, **kwargs):
            raise RuntimeError("Milvus unavailable at creation")
        patch.setattr(service.store, "ensure_collection", unavailable)
        with pytest.raises(RuntimeError, match="Milvus unavailable"):
            service.create("unfinished S2 creation")
    with metadata.connect() as db:
        failed = db.execute("SELECT id FROM knowledge_bases WHERE name='unfinished S2 creation'").fetchone()[0]
        # Reconstruct the old column layout; lifecycle state comes from service operations.
        if version == 1:
            db.execute("ALTER TABLE chunks DROP COLUMN original_path")
            db.execute("ALTER TABLE documents DROP COLUMN removed")
            db.execute("ALTER TABLE knowledge_bases DROP COLUMN pending_operation")
        db.execute("ALTER TABLE documents DROP COLUMN updated_at")
        db.execute(f"PRAGMA user_version = {version}")
        previous = {table: [dict(row) for row in db.execute(f"SELECT * FROM {table} ORDER BY id")]
                    for table in ("knowledge_bases", "documents", "chunks")}
    Metadata(service.root)
    with metadata.connect() as db:
        assert db.execute("PRAGMA user_version").fetchone()[0] == 3
        for table, rows in previous.items():
            current = [dict(row) for row in db.execute(f"SELECT * FROM {table} ORDER BY id")]
            assert [{key: row[key] for key in rows[0]} for row in current] == rows
    assert service.metadata.document(document["id"])["updated_at"] is None
    assert all(kb["updated_at"] is None for kb in service.list_libraries())
    state = service.status(kb)
    Metadata(service.root)  # Reopening v3 is idempotent.
    assert service.status(kb) == state
    assert service.source(kb, hit.chunk_id)["original_path"] == document["original_path"]
    assert service.search(kb, "test").hits[0] == hit
    assert service.import_document(kb, source)["unchanged"]
    assert service.status(failed)["pending_operation"] == "create"
    assert service.retry(failed)["state"] == "READY"
    with metadata.connect() as db:
        assert {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")} == {"knowledge_bases", "documents", "chunks"}
        db.execute("PRAGMA user_version = 99")
    with pytest.raises(ValueError, match="schema version: 99"):
        Metadata(service.root)


def test_library_list_is_local_read_only_and_includes_unfinished_bases(service, tmp_path):
    assert service.list_libraries() == []
    populated = service.create("Zulu")["id"]
    empty = service.create("Alpha")["id"]
    source = tmp_path / "source.md"
    source.write_text("A document", encoding="utf-8")
    document = service.import_document(populated, source)
    source = tmp_path / "removed.md"
    source.write_text("A removed document", encoding="utf-8")
    removed = service.remove(populated, service.import_document(populated, source)["id"])
    # Registered creation targets can be listed even before their profile exists.
    service.metadata.register("a" * 32, "Pending", "pending_collection", service.profile_hash)
    service.metadata.register("b" * 32, "Repair", "repair_collection", service.profile_hash)
    service.metadata.set_state("b" * 32, "NEEDS_REPAIR", "creation failed")
    offline = KnowledgeService(replace(service.config, embedding_revision="unavailable-profile"))
    before = service.metadata.path.read_bytes()
    try:
        libraries = offline.list_libraries()
        assert [kb["name"] for kb in libraries] == ["Alpha", "Pending", "Repair", "Zulu"]
        assert libraries == offline.metadata.list_libraries() == offline.list_libraries()
        by_id = {kb["id"]: kb for kb in libraries}
        assert by_id[empty]["document_count"] == 0 and by_id[empty]["updated_at"] is None
        assert by_id[populated]["document_count"] == 1
        assert by_id[populated]["updated_at"] == removed["updated_at"]
        assert by_id[populated]["revision"] == 3
        assert by_id["a" * 32]["state"] == "UPDATING"
        assert by_id["b" * 32]["state"] == "NEEDS_REPAIR"
        assert by_id["b" * 32]["error"] == "creation failed"
        for library in libraries:
            assert {key: library[key] for key in service.metadata.library(library["id"])} == service.metadata.library(library["id"])
        assert offline._store is None and not offline.embedding.loaded
        assert offline.metadata.document(document["id"])["updated_at"] == document["updated_at"]
        assert service.metadata.path.read_bytes() == before
    finally:
        offline.close()


def test_document_update_time_tracks_successful_commits_and_retry(service, tmp_path, monkeypatch):
    class Clock:
        value = datetime(2026, 1, 1, tzinfo=timezone.utc)

        @classmethod
        def now(cls, tz):
            assert tz == timezone.utc
            return cls.value

    monkeypatch.setattr("codeplus.knowledge.metadata.datetime", Clock)
    kb = service.create("timestamps")["id"]
    source = tmp_path / "source.md"
    source.write_text("Original document", encoding="utf-8")
    original = service.import_document(kb, source)
    assert original["updated_at"] == Clock.value.isoformat()
    Clock.value += timedelta(days=1)
    assert service.import_document(kb, source)["updated_at"] == original["updated_at"]
    source.write_text("Replacement document", encoding="utf-8")
    with monkeypatch.context() as patch:
        def fail_verify(*args):
            raise RuntimeError("verify failed")
        patch.setattr(service.store, "verify_document", fail_verify)
        with pytest.raises(RuntimeError, match="verify failed"):
            service.import_document(kb, source)
    assert service.metadata.document(original["id"])["updated_at"] == original["updated_at"]
    assert service.list_libraries()[0]["updated_at"] == original["updated_at"]
    Clock.value += timedelta(days=1)
    replaced = service.retry(kb)
    assert replaced["updated_at"] == Clock.value.isoformat()
    assert replaced["id"] == original["id"] and replaced["generation_id"] != original["generation_id"]
    Clock.value += timedelta(days=1)
    # Fail after the document timestamp UPDATE, proving it rolls back with revision.
    with service.metadata.connect() as db:
        db.execute("CREATE TRIGGER fail_revision BEFORE UPDATE OF revision ON knowledge_bases "
                   "BEGIN SELECT RAISE(ABORT, 'revision failed'); END")
    with pytest.raises(sqlite3.IntegrityError, match="revision failed"):
        service.remove(kb, original["id"])
    failed = service.metadata.document(original["id"])
    assert failed["updated_at"] == replaced["updated_at"] and not failed["removed"]
    assert service.status(kb)["revision"] == 2
    with service.metadata.connect() as db:
        db.execute("DROP TRIGGER fail_revision")
    Clock.value += timedelta(days=1)
    removed = service.retry(kb)
    assert removed["removed"] and removed["updated_at"] == Clock.value.isoformat()
    assert service.list_libraries()[0]["updated_at"] == removed["updated_at"]
    assert service.list_libraries()[0]["document_count"] == 0
    Clock.value += timedelta(days=1)
    assert service.remove(kb, original["id"])["updated_at"] == removed["updated_at"]
    assert service.retry(kb)["unchanged"]
    assert service.metadata.document(original["id"])["updated_at"] == removed["updated_at"]
    source.write_text("Original document", encoding="utf-8")
    restored = service.import_document(kb, source)
    assert restored["updated_at"] == Clock.value.isoformat()
    assert restored["generation_id"] == original["generation_id"]
    assert restored["original_path"] == original["original_path"]
    assert service.status(kb)["revision"] == 4


def test_process_lock_crash_keeps_import_pending(service, tmp_path):
    kb = service.create("writer")["id"]
    other = service.create("independent")["id"]
    source = tmp_path / "process.md"
    source.write_text("# Process\n\nImported in a different process", encoding="utf-8")
    marker = tmp_path / "writing"
    command = [sys.executable, "-c", """
import json, os, sys
from pathlib import Path
sys.path.insert(0, str(Path.cwd()))
sys.path.insert(0, str(Path.cwd() / 'tests'))
import codeplus
assert Path(codeplus.__file__).resolve().parent.parent == Path.cwd()
from test_knowledge_service import TinyEmbedding, MemoryStore
from codeplus.config import KnowledgeConfig
from codeplus.knowledge.service import KnowledgeService
service = KnowledgeService(KnowledgeConfig(**json.loads(sys.argv[1])))
service.embedding = TinyEmbedding()
class PausedStore(MemoryStore):
    def upsert_chunks(self, name, rows):
        Path(sys.argv[4]).touch()
        sys.stdin.readline()
        os._exit(0)
service._store = PausedStore()
with service._locked(sys.argv[2]) as bound:
    service._collection(bound, create=True)
service.import_document(sys.argv[2], sys.argv[3])
""", json.dumps(asdict(service.config)), kb, str(source), str(marker)]
    writer = subprocess.Popen(command, stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    reader = None
    try:
        deadline = time.monotonic() + 15
        while not marker.exists() and time.monotonic() < deadline and writer.poll() is None:
            time.sleep(0.05)
        assert marker.exists(), "writer did not reach pending SDK write"
        assert service.status(other)["state"] == "READY"
        reader = subprocess.Popen([sys.executable, "-c", """
import json, sys
from pathlib import Path
import codeplus
assert Path(codeplus.__file__).resolve().parent.parent == Path.cwd()
from codeplus.config import KnowledgeConfig
from codeplus.knowledge.service import KnowledgeService
service = KnowledgeService(KnowledgeConfig(**json.loads(sys.argv[1])))
print('waiting', flush=True)
try:
    service.search(sys.argv[2], 'test')
except ValueError as exc:
    assert 'UPDATING' in str(exc), str(exc)
    print('blocked by pending state', flush=True)
else:
    raise AssertionError('incomplete import was searchable')
""", json.dumps(asdict(service.config)), kb], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        assert reader.stdout.readline().strip() == "waiting"
        time.sleep(0.2)
        assert reader.poll() is None, "search bypassed the writer's OS lock"
        writer.kill()  # A real process death: no finally/SQLite success/status repair runs.
        writer.communicate(timeout=10)
        output, error = reader.communicate(timeout=15)
        assert reader.returncode == 0 and "blocked by pending state" in output, error
        state = service.status(kb)
        assert state["state"] == "UPDATING" and state["revision"] == 0
        assert Path(state["documents"][0]["pending_path"]).exists()
    finally:
        for process in (writer, reader):
            if process is not None and process.poll() is None:
                process.kill()
                process.communicate(timeout=10)


def write_pdf(path, pages, *, encrypted=False):
    """Small real PDF fixture, without adding a second PDF library to test dependencies."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(300, 200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("test-password")
    writer.write(path)


def test_pdf_docx_locations_and_format_errors(service, tmp_path):
    from docx import Document
    from codeplus.knowledge.documents import parse_document
    from codeplus.knowledge.citations import source_location

    kb = service.create("formats")["id"]
    pdf = tmp_path / "two pages.pdf"
    write_pdf(pdf, ["Amber license expires in April.", "Cobalt delivery arrives in November."])
    doc = service.import_document(kb, pdf)
    hits = service.search(kb, "Cobalt", 50).hits
    assert {span.page for hit in hits for span in hit.source_spans} == {1, 2}
    assert "第2页" in source_location([asdict(s) for h in hits for s in h.source_spans])
    text, _ = parse_document(Path(doc["original_path"]))
    assert any("Cobalt" in text[s.char_start:s.char_end] and s.page == 2 for h in hits for s in h.source_spans)
    word = tmp_path / "body order.docx"
    source = Document()
    source.add_heading("Dispatch", 1)
    source.add_paragraph("Before table")
    table = source.add_table(rows=2, cols=2)
    for cell, value in zip((c for r in table.rows for c in r.cells), ("Route", "Deadline", "Cobalt", "November")):
        cell.text = value
    source.add_paragraph("After table")
    source.save(word)
    document = service.import_document(kb, word)
    text, blocks = parse_document(Path(document["original_path"]))
    assert text.index("Before table") < text.index("November") < text.index("After table")
    assert next(b.source for b in blocks if "After table" in b.text).paragraph == 3
    hits = [h for h in service.search(kb, "November", 50).hits if h.doc_id == document["id"]]
    spans = [s for h in hits for s in h.source_spans]
    assert all(s.page is None and s.line_start is None and s.heading_path == ["Dispatch"] for s in spans)
    assert any((s.table, s.row, s.column) == (1, 2, 2) and "November" in text[s.char_start:s.char_end] for s in spans)
    assert all(text[s.char_start:s.char_end] in h.text for h in hits for s in h.source_spans)
    label = source_location([asdict(s) for s in spans])
    assert "第3段" in label and "表1 · 第2行 · 第2列" in label
    before = service.status(kb)
    invalid = tmp_path / "invalid.pdf"
    write_pdf(invalid, [""])
    with pytest.raises(ValueError, match="empty or scanned"):
        service.import_document(kb, invalid)
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject
    writer = PdfWriter()
    page = writer.add_blank_page(300, 200)
    image = DecodedStreamObject()
    image.set_data(b"\x00")
    image.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Image"),
                  NameObject("/Width"): NumberObject(1), NameObject("/Height"): NumberObject(1),
                  NameObject("/ColorSpace"): NameObject("/DeviceGray"), NameObject("/BitsPerComponent"): NumberObject(8)})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Im1"): writer._add_object(image)})})
    stream = DecodedStreamObject()
    stream.set_data(b"q 100 0 0 100 0 0 cm /Im1 Do Q")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(invalid)
    with pytest.raises(ValueError, match="empty or scanned"):
        service.import_document(kb, invalid)
    write_pdf(invalid, ["Secret"], encrypted=True)
    with pytest.raises(ValueError, match="Encrypted PDF"):
        service.import_document(kb, invalid)
    invalid.write_bytes(b"not a PDF")
    with pytest.raises(ValueError, match="Cannot parse PDF"):
        service.import_document(kb, invalid)
    with pytest.raises(ValueError, match="convert legacy .doc"):
        service.import_document(kb, tmp_path / "old.doc")
    assert service.status(kb) == before


@pytest.mark.skipif(not os.getenv("CODEPLUS_TEST_MILVUS_URI"), reason="real Milvus not requested")
@pytest.mark.parametrize("legacy", [False, True], ids=["hybrid", "legacy-dense"])
def test_real_daily_retrieval_modes(service, tmp_path, monkeypatch, legacy):
    from codeplus.knowledge.milvus_store import MilvusStore

    service._store = MilvusStore(os.environ["CODEPLUS_TEST_MILVUS_URI"])
    if legacy:
        service.profile.pop("indexing")
        service.profile_hash = fingerprint(service.profile)
    try:
        kb = service.create("daily retrieval")["id"]
        service.profile = profile(service.config)
        service.profile_hash = fingerprint(service.profile)
        documents = []
        for name, text, vector in (("semantic", "Semantic evidence", [1.0, 0.0, 0.0]),
                                   ("keyword", "红杉定额 ZXQ741", [0.0, 1.0, 0.0])):
            source = tmp_path / f"{name}.md"
            source.write_text(text, encoding="utf-8")
            monkeypatch.setattr(service.embedding, "encode_documents", lambda texts, vector=vector: [vector for _ in texts])
            documents.append(service.import_document(kb, source))
        service.config = replace(service.config, retrieval_mode="dense")
        dense = service.search(kb, "ZXQ741", 1)
        assert dense.hits[0].doc_id == documents[0]["id"] and dense.hits[0].score_type == "cosine_similarity"
        service.config = replace(service.config, retrieval_mode="bm25")
        with monkeypatch.context() as patch:
            patch.setattr(service.embedding, "encode_query", lambda *_: pytest.fail("BM25 encoded a query"))
            if legacy:
                with pytest.raises(ValueError, match="create a new hybrid base"):
                    service.search(kb, "ZXQ741", 1)
            else:
                bm25 = service.search(kb, "ZXQ741", 1)
                assert bm25.hits[0].doc_id == documents[1]["id"] and bm25.hits[0].score_type == "bm25"
                assert bm25.hits[0].score > 0 and bm25.retrieval["lane_counts"] == {"dense": None, "bm25": 1}
                chinese = service.search(kb, "红杉", 1)
                assert chinese.hits[0].chunk_id == bm25.hits[0].chunk_id
        service.config = replace(service.config, retrieval_mode="auto", retrieval_candidates=2)
        result = service.search(kb, "ZXQ741", 2)
        assert result.retrieval["mode"] == ("dense" if legacy else "hybrid")
        assert result.retrieval["lane_counts"] == {"dense": 2, "bm25": None if legacy else 1}
        assert {hit.doc_id for hit in result.hits} == {doc["id"] for doc in documents}
        if not legacy:
            assert result.hits[0].doc_id == documents[1]["id"]
            assert [hit.score for hit in result.hits] == pytest.approx([1 / 61 + 1 / 62, 1 / 61])
            assert all(hit.score_type == "rrf" for hit in result.hits)
            service.config = replace(service.config, retrieval_mode="hybrid")
            assert service.search(kb, "ZXQ741", 2).hits == result.hits
        for hit in result.hits:
            saved = service.source(kb, hit.chunk_id)
            assert saved["source_status"] == "current" and saved["generation_id"] == hit.generation_id
            assert saved["text"] == hit.text
        print(f"daily retrieval legacy={legacy}: dense/BM25/auto modes and sources verified")
    finally:
        for owned in service.list_libraries():
            name = owned["collection_name"]
            if service.store.client.has_collection(name):
                service.store.client.drop_collection(name)
            assert not service.store.client.has_collection(name)


@pytest.mark.skipif(not os.getenv("CODEPLUS_TEST_MILVUS_URI"), reason="real Milvus not requested")
@pytest.mark.parametrize("legacy", [False, True], ids=["hybrid", "legacy-dense"])
def test_real_store_document_isolation_and_binding(service, tmp_path, monkeypatch, legacy):
    from codeplus.knowledge.milvus_store import MilvusStore
    from codeplus.knowledge.models import Chunk

    chunk = Chunk("chunk", "doc", "generation", 0, "中" * 22000, [])
    with pytest.raises(ValueError, match="UTF-8 bytes"):
        MilvusStore.rows([chunk], [[1.0, 0.0, 0.0]], 3)
    with pytest.raises(ValueError, match="dimension"):
        MilvusStore.rows([replace(chunk, text="small")], [[1.0, 0.0]], 3)
    if legacy:
        service.profile.pop("indexing")
        service.profile_hash = fingerprint(service.profile)
    service._store = MilvusStore(os.environ["CODEPLUS_TEST_MILVUS_URI"])
    service.config = replace(service.config, milvus_uri=os.environ["CODEPLUS_TEST_MILVUS_URI"])
    kb = None
    try:
        # Dense creation is interrupted before its index; hybrid between its two indexes.
        with monkeypatch.context() as patch:
            create_index = service.store.client.create_index
            def fail_index(*args, **kwargs):
                if legacy or service.store.client.list_indexes(args[0]):
                    raise RuntimeError("before index creation")
                return create_index(*args, **kwargs)
            patch.setattr(service.store.client, "create_index", fail_index)
            with pytest.raises(RuntimeError, match="before index creation"):
                service.create("real SDK")
        with service.metadata.connect() as db:
            kb = dict(db.execute("SELECT * FROM knowledge_bases WHERE name='real SDK'").fetchone())
        assert service.store.client.has_collection(kb["collection_name"])
        assert service.store.client.list_indexes(kb["collection_name"]) == ([] if legacy else ["dense"])
        # The latest service must recover using the collection's saved contract.
        service.profile = profile(service.config)
        service.profile_hash = fingerprint(service.profile)
        assert service.retry(kb["id"])["state"] == "READY"
        assert set(service.store.client.list_indexes(kb["collection_name"])) == ({"dense"} if legacy else {"dense", "sparse"})
        docs = []
        for name in ("first", "second"):
            source = tmp_path / f"{name}.md"
            source.write_text(f"# {name}\n\nA passage for {name}", encoding="utf-8")
            docs.append(service.import_document(kb["id"], source))
        for document in docs:
            assert document["chunk_count"] == 1
        result = service.search(kb["id"], "test")
        assert result.retrieval["mode"] == ("dense" if legacy else "hybrid")
        assert all(hit.score_type == ("cosine_similarity" if legacy else "rrf") for hit in result.hits)
        assert {hit.doc_id for hit in result.hits} == {doc["id"] for doc in docs}
        if not legacy:
            assert len(service.store.search_bm25(kb["collection_name"], "passage", 10)) == 2
        fields = ["chunk_id", "doc_id", "generation_id", "text", "dense"]
        rows = list(service.store.client.query(kb["collection_name"], filter="", limit=10, output_fields=fields))
        service.store.upsert_chunks(kb["collection_name"], rows)
        for doc in docs:
            service.store.verify_document(kb["collection_name"], doc["id"], [r for r in rows if r["doc_id"] == doc["id"]])
        with pytest.raises(ValueError, match="profile mismatch"):
            service.store.ensure_collection(kb["collection_name"], "different", 3)
        with pytest.raises(ValueError, match="profile mismatch"):
            service.store.ensure_collection(kb["collection_name"], kb["profile_hash"], 4)
        child = tmp_path / "crash.py"
        child.write_text('''
import json, os, sys
from pathlib import Path
sys.path.insert(0, str(Path.cwd()))
sys.path.insert(0, str(Path.cwd() / 'tests'))
import codeplus
assert Path(codeplus.__file__).resolve().parent.parent == Path.cwd()
from test_knowledge_service import TinyEmbedding
from codeplus.config import KnowledgeConfig
from codeplus.knowledge.service import KnowledgeService
service = KnowledgeService(KnowledgeConfig(**json.loads(sys.argv[1])))
service.embedding = TinyEmbedding()
kb, path, doc, phase = sys.argv[2:]
if phase in ('deleted', 'remove'):
    erase = service.store.delete_document
    def stop(*args):
        erase(*args)
        os._exit(73)
    service.store.delete_document = stop
elif phase == 'partial':
    write = service.store.upsert_chunks
    def stop(name, rows):
        assert len(rows) > 1
        write(name, rows[:len(rows)//2])
        os._exit(73)
    service.store.upsert_chunks = stop
else:
    service.metadata.finish_import = lambda *args, **kwargs: os._exit(73)
if phase == 'remove':
    service.remove(kb, doc)
else:
    service.import_document(kb, path)
''', encoding="utf-8")
        source = tmp_path / "first.md"
        for phase in ("deleted", "partial", "written", "remove"):
            before = service.status(kb["id"])
            old = service.metadata.document(docs[0]["id"])
            source.write_text(f"# {phase}\n\n" + "New content with distinct chunks. " * 8, encoding="utf-8")
            process = subprocess.run([sys.executable, str(child), json.dumps(asdict(service.config)), kb["id"],
                                      str(source), docs[0]["id"], phase], capture_output=True, text=True, timeout=90)
            assert process.returncode == 73, process.stderr
            with KnowledgeService(service.config).metadata.connect() as db:
                assert db.execute("SELECT generation_id FROM documents WHERE id=?", (old["id"],)).fetchone()[0] == old["generation_id"]
            reopened = KnowledgeService(service.config)
            reopened._store = service.store
            reopened.embedding = TinyEmbedding()
            assert reopened.status(kb["id"])["state"] == "UPDATING"
            with pytest.raises(ValueError, match="UPDATING"):
                reopened.search(kb["id"], "test")
            pending_doc = reopened.metadata.document(old["id"])
            expected = [] if phase == "remove" else json.loads(Path(pending_doc["pending_path"]).read_text(encoding="utf-8"))["rows"]
            recovered = reopened.retry(kb["id"])
            state = reopened.status(kb["id"])
            assert state["revision"] == before["revision"] + 1 and state["state"] == "READY"
            assert state["collection_name"] == kb["collection_name"]
            service.store.verify_document(kb["collection_name"], old["id"], expected)
            actual_b = service.store.client.query(kb["collection_name"], filter="doc_id == {doc_id}",
                                                  filter_params={"doc_id": docs[1]["id"]}, output_fields=fields)
            assert list(actual_b) == [r for r in rows if r["doc_id"] == docs[1]["id"]]
            if not legacy:
                sparse_hits = service.store.search_bm25(kb["collection_name"], "distinct", 50)
                assert {h["chunk_id"] for h in sparse_hits} <= {r["chunk_id"] for r in expected}
                assert bool(sparse_hits) == (phase != "remove")
            assert recovered["removed"] == (phase == "remove")
            assert reopened.retry(kb["id"])["unchanged"]
            print(f"real process exit at {phase}: blocked then READY revision={state['revision']}; B unchanged")
        if not legacy:
            # Check the server's actual BM25 parameters, including flattened numeric strings.
            name = kb["collection_name"]
            service.store.client.release_collection(name)
            service.store.client.drop_index(name, "sparse")
            wrong = deepcopy(service.profile["indexing"]["sparse"])
            wrong["params"]["bm25_k1"] = 1.8
            indexes = service.store.client.prepare_index_params()
            indexes.add_index("sparse", **wrong)
            service.store.client.create_index(name, indexes)
            with service._locked(kb["id"]) as bound, pytest.raises(ValueError, match="index sparse"):
                service._collection(bound)
    finally:
        for owned in service.list_libraries():
            name = owned["collection_name"]
            if service.store.client.has_collection(name):
                service.store.client.drop_collection(name)
            assert not service.store.client.has_collection(name)
